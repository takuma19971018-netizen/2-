"""共有スプレッドシート「◆英語分析◆」(xlsx書き出し) を年月確定済みJSONに変換する。

シートの1行目は年(2011, =B1+1, =F1+1 ...)、2行目は月(2/4/7/10)。
年は数式なので、参照先を辿って実際の年に解決する。

使い方:
    pip install openpyxl
    python scripts/sheet_to_json.py eigo.xlsx data/sheet
"""
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string


def resolve_years(row1):
    def ev(v):
        if isinstance(v, (int, float)):
            return int(v)
        m = re.match(r"=([A-Z]+)1\+1", str(v))
        return ev(row1[column_index_from_string(m.group(1)) - 1]) + 1

    years, cur = [], None
    for v in row1:
        if v not in (None, ""):
            cur = ev(v)
        years.append(cur)
    return years


def main(xlsx, outdir):
    wb = openpyxl.load_workbook(xlsx)
    for ws in wb:
        years = resolve_years([c.value for c in ws[1]])
        months = [c.value for c in ws[2]]
        rows = []
        for i, r in enumerate(ws.iter_rows(min_row=3), start=3):
            title = r[0].value
            marks = []
            for j, c in enumerate(r[1:], 1):
                if c.value in (None, ""):
                    continue
                note = "" if c.value == "●" else " " + str(c.value)
                marks.append(f"{years[j]}/{int(months[j]):02d}{note}")
            if title or marks:
                rows.append({"row": i, "title": title, "marks": marks})
        name = ws.title.strip().replace("Ｎ", "N")
        Path(outdir, f"sheet_{name}_raw.json").write_text(
            json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
